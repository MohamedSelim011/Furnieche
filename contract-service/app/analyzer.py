import base64
import io
import json
import logging
from dataclasses import dataclass
from typing import Any
from urllib.parse import urlparse

import anthropic
import httpx
from docx import Document

from .config import Settings
from .schema import ContractAnalysis, output_schema

log = logging.getLogger(__name__)

SYSTEM_PROMPT = """You analyze construction, interior fit-out, and furniture contracts for Furniche, \
an app used by engineers and contractors in Egypt and the wider MENA region to run furnishing projects.

The person reading your analysis is the engineer/contractor who uploaded the contract. Extract what the \
document actually says, and point out what could cost them time or money.

Rules:
- Only report terms that appear in the document. If something is not stated, use null (or an empty list). \
Never invent parties, dates, amounts, or clauses.
- Contracts may be in Arabic, English, or both. Read both. Write every free-text field in Arabic if the \
contract is mainly Arabic, otherwise in English, and set output_language accordingly. Keep proper names as written.
- Normalize dates to YYYY-MM-DD (convert Arabic-Indic digits and Hijri dates when the Gregorian date is \
unambiguous). Normalize money to plain numbers and put the currency in financials.currency.
- For payment milestones given only as a percentage, compute the amount from the total value when the total is stated.
- Use the clause numbering exactly as it appears in the document for references.
- Risks: focus on what matters to the contractor — uncapped delay penalties, vague scope, payment tied to \
subjective approval, missing variation-order process, one-sided termination, unclear warranty start, \
missing price-escalation terms, and similar. Rank by severity.
- If the document is not a contract (e.g. an invoice or a drawing), set is_contract to false and fill \
what you can."""

USER_INSTRUCTION = "Analyze this contract and return the structured analysis."

IMAGE_TYPES = {"image/jpeg", "image/png", "image/webp", "image/gif"}
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
EXTENSION_TYPES = {
    "pdf": "application/pdf",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "png": "image/png",
    "webp": "image/webp",
    "gif": "image/gif",
    "docx": DOCX_TYPE,
    "txt": "text/plain",
}


class AnalysisError(Exception):
    """A failure worth showing to the user as-is."""


@dataclass
class AnalysisOutcome:
    result: dict[str, Any]
    model: str
    input_tokens: int
    output_tokens: int


def _media_type(file_name: str, header_type: str | None) -> str:
    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if ext in EXTENSION_TYPES:
        return EXTENSION_TYPES[ext]
    return (header_type or "").split(";")[0].strip().lower()


def download(url: str, file_name: str, settings: Settings) -> tuple[bytes, str]:
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname != settings.allowed_file_host:
        raise AnalysisError("File URL is not on the allowed storage host")

    with httpx.Client(timeout=60.0, follow_redirects=False) as http:
        with http.stream("GET", url) as res:
            if res.status_code != 200:
                raise AnalysisError(f"Could not download the file (HTTP {res.status_code})")
            chunks: list[bytes] = []
            size = 0
            for chunk in res.iter_bytes():
                size += len(chunk)
                if size > settings.max_file_bytes:
                    raise AnalysisError("File is too large to analyze")
                chunks.append(chunk)
            return b"".join(chunks), _media_type(file_name, res.headers.get("content-type"))


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def build_content(data: bytes, media_type: str) -> list[dict[str, Any]]:
    """The document goes before the instruction, as the API recommends."""
    if media_type == "application/pdf":
        block: dict[str, Any] = {
            "type": "document",
            "source": {"type": "base64", "media_type": "application/pdf", "data": base64.b64encode(data).decode()},
        }
    elif media_type in IMAGE_TYPES:
        block = {
            "type": "image",
            "source": {"type": "base64", "media_type": media_type, "data": base64.b64encode(data).decode()},
        }
    elif media_type in (DOCX_TYPE, "text/plain"):
        text = _docx_text(data) if media_type == DOCX_TYPE else data.decode("utf-8", errors="replace")
        if not text.strip():
            raise AnalysisError("The document contains no readable text")
        block = {
            "type": "document",
            "source": {"type": "text", "media_type": "text/plain", "data": text},
        }
    else:
        raise AnalysisError("Unsupported file type. Upload a PDF, Word (.docx), or photo of the contract.")
    return [block, {"type": "text", "text": USER_INSTRUCTION}]


def analyze(client: anthropic.Anthropic, url: str, file_name: str, settings: Settings) -> AnalysisOutcome:
    data, media_type = download(url, file_name, settings)
    content = build_content(data, media_type)

    try:
        with client.beta.messages.stream(
            model=settings.model,
            max_tokens=32000,
            # Re-runs the request on Anthropic's recommended model if a safety
            # classifier declines it, instead of returning a refusal.
            betas=["server-side-fallback-2026-07-01"],
            fallbacks="default",
            thinking={"type": "adaptive"},
            output_config={
                "effort": "high",
                "format": {"type": "json_schema", "schema": output_schema()},
            },
            system=SYSTEM_PROMPT,
            messages=[{"role": "user", "content": content}],
        ) as stream:
            message = stream.get_final_message()
    except anthropic.BadRequestError as e:
        log.error("Claude rejected the request: %s", e.message)
        # Covers both bad documents and account problems (e.g. no API credit);
        # the exact reason is in the service log.
        raise AnalysisError(
            "The AI service rejected the request. Check the API account's billing, or try a different file."
        ) from e
    except anthropic.RateLimitError as e:
        raise AnalysisError("The AI service is busy. Try again in a minute.") from e
    except anthropic.APIStatusError as e:
        log.error("Claude API error %s: %s", e.status_code, e.message)
        raise AnalysisError("The AI service returned an error. Try again later.") from e
    except anthropic.APIConnectionError as e:
        raise AnalysisError("Could not reach the AI service") from e

    if message.stop_reason == "refusal":
        raise AnalysisError("The AI declined to analyze this document")
    if message.stop_reason == "max_tokens":
        raise AnalysisError("The contract is too long to analyze in one pass")

    text = next((b.text for b in reversed(message.content) if b.type == "text"), None)
    if text is None:
        raise AnalysisError("The AI returned no analysis")

    try:
        parsed = ContractAnalysis.model_validate(json.loads(text))
    except (json.JSONDecodeError, ValueError) as e:
        log.error("Invalid analysis JSON: %s", e)
        raise AnalysisError("The AI returned an invalid analysis. Try again.") from e

    return AnalysisOutcome(
        result=parsed.model_dump(),
        model=message.model,
        input_tokens=message.usage.input_tokens,
        output_tokens=message.usage.output_tokens,
    )
